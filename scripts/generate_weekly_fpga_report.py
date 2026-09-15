from __future__ import annotations

import json
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\fpga")
TEMPLATE = Path(r"C:\Users\DELL\Downloads\report2.docx")
RESULT_ROOT = ROOT / "final_results" / "zybo_z7_10_qnn_60fps_104mhz"
OUT_DIR = ROOT / "reports" / "weekly_fpga_2026_08_25"
OUT_DOCX = OUT_DIR / "Bao_cao_tuan_FPGA_Zybo_Z7_10_2026_08_25.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color=None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_paragraph(doc, text: str = "", *, style: str = "Normal", align=None, first_line=True, space_after=4):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.25
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(0.75)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, 12)
    return paragraph


def add_heading(doc, text: str, level: int) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, 15 if level == 1 else 13, True, (31, 78, 121))
    return paragraph


def add_caption(doc, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, 10, False, (80, 80, 80))
    run.italic = True


def make_plots(results: dict, video: dict) -> tuple[Path, Path, Path]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    plt.rcParams["font.family"] = "DejaVu Sans"

    fps_chart = OUT_DIR / "fps_comparison.png"
    labels = ["Baseline\n100 MHz", "Bản tối ưu\n104 MHz"]
    values = [
        results["baseline_100mhz"]["dual_roi_frames_per_second"],
        results["hardware_benchmark"]["dual_roi_frames_per_second"],
    ]
    fig, ax = plt.subplots(figsize=(7.2, 4.2), dpi=180)
    bars = ax.bar(labels, values, color=["#7aa6d8", "#1f77b4"], width=0.55)
    ax.axhline(60, color="#c0392b", linewidth=1.5, linestyle="--", label="Mục tiêu 60 FPS")
    ax.set_ylim(0, 70)
    ax.set_ylabel("Khung hình/s cho hai ROI")
    ax.set_title("Thông lượng QNN đo trực tiếp trên FPGA")
    ax.legend(loc="lower right")
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 1.2, f"{value:.2f}", ha="center", fontsize=10, fontweight="bold")
    fig.tight_layout()
    fig.savefig(fps_chart, bbox_inches="tight")
    plt.close(fig)

    resource_chart = OUT_DIR / "resource_utilization.png"
    resource = results["resources"]
    labels = ["LUT", "FF", "BRAM", "DSP"]
    values = [resource["slice_luts_percent"], resource["slice_registers_percent"], resource["bram_percent"], resource["dsp_percent"]]
    fig, ax = plt.subplots(figsize=(7.2, 4.2), dpi=180)
    colors = ["#e67e22", "#3498db", "#27ae60", "#8e44ad"]
    bars = ax.bar(labels, values, color=colors, width=0.58)
    ax.set_ylim(0, 100)
    ax.set_ylabel("Mức sử dụng (%)")
    ax.set_title("Tài nguyên Zybo Z7-10 sau implementation")
    ax.grid(axis="y", linewidth=0.4, alpha=0.45)
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 2, f"{value:.2f}%", ha="center", fontsize=10, fontweight="bold")
    fig.tight_layout()
    fig.savefig(resource_chart, bbox_inches="tight")
    plt.close(fig)

    architecture = OUT_DIR / "architecture.png"
    fig, ax = plt.subplots(figsize=(11, 3.8), dpi=180)
    ax.axis("off")
    boxes = [
        (0.03, 0.48, "Video\nkính hiển vi"),
        (0.20, 0.48, "Hai ROI\n120 x 120"),
        (0.37, 0.48, "QNN W4A6\n96 x 96"),
        (0.54, 0.48, "Tracker ID\n+ đối chiếu"),
        (0.71, 0.48, "Đếm một lần\nROI 1 -> ROI 2"),
        (0.88, 0.48, "Kết quả\ncell / droplet"),
    ]
    colors = ["#f7dc6f", "#aed6f1", "#a9dfbf", "#d7bde2", "#f5cba7", "#f1948a"]
    for index, ((x, y, text), color) in enumerate(zip(boxes, colors)):
        rect = plt.Rectangle((x, y - 0.14), 0.12, 0.28, facecolor=color, edgecolor="#1f4e79", linewidth=1.2)
        ax.add_patch(rect)
        ax.text(x + 0.06, y, text, ha="center", va="center", fontsize=9, fontweight="bold")
        if index < len(boxes) - 1:
            ax.annotate("", xy=(x + 0.16, y), xytext=(x + 0.122, y), arrowprops=dict(arrowstyle="->", color="#1f1f1f", lw=1.3))
    ax.text(0.5, 0.9, "Luồng xử lý QNN hai ROI chạy trên Zybo Z7-10", ha="center", va="center", fontsize=14, fontweight="bold", color="#1f4e79")
    ax.text(0.5, 0.12, "ROI 2 chỉ xác nhận đối tượng của ROI 1; không cộng hai lần.", ha="center", va="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(architecture, bbox_inches="tight", transparent=False)
    plt.close(fig)
    return fps_chart, resource_chart, architecture


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_metric_table(doc, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Cm(7.2)
    table.columns[1].width = Cm(8.0)
    hdr = table.rows[0].cells
    hdr[0].text = "Thông số"
    hdr[1].text = "Kết quả"
    for cell in hdr:
        set_cell_shading(cell, "1F4E79")
        for run in cell.paragraphs[0].runs:
            set_run_font(run, 11, True, (255, 255, 255))
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_repeat_table_header(table.rows[0])
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
        for cell in cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                for run in paragraph.runs:
                    set_run_font(run, 10.5)
    doc.add_paragraph()


def main() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    results = json.loads((RESULT_ROOT / "results.json").read_text(encoding="utf-8"))
    video = json.loads((RESULT_ROOT / "video_3_4_exact_tracking" / "report.json").read_text(encoding="utf-8"))
    fps_chart, resource_chart, architecture = make_plots(results, video)
    preview = RESULT_ROOT / "video_3_4_exact_tracking" / "preview.jpg"

    doc = Document(TEMPLATE)
    clear_body(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(64)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("BÁO CÁO TIẾN ĐỘ TUẦN")
    set_run_font(run, 22, True, (31, 78, 121))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    run = subtitle.add_run("HỆ THỐNG PHÁT HIỆN GIỌT VÀ TẾ BÀO TRÊN FPGA")
    set_run_font(run, 17, True, (31, 78, 121))

    for text in [
        "Triển khai QNN hai ROI trên Zybo Z7-10",
        "Thời gian báo cáo: tuần kết thúc ngày 25/08/2026",
        "Sinh viên thực hiện: Nhật Minh Dương",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(text)
        set_run_font(r, 13 if text.startswith("Triển") else 12)

    doc.add_page_break()
    toc = doc.add_paragraph()
    toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = toc.add_run("MỤC LỤC")
    set_run_font(r, 18, True, (31, 78, 121))
    toc.paragraph_format.space_after = Pt(16)
    for entry in [
        "1. Mục tiêu tuần",
        "2. Công việc đã thực hiện",
        "3. Kết quả triển khai QNN trên FPGA",
        "4. Kiểm chứng video và thuật toán đếm hai ROI",
        "5. Tích hợp camera Phantom Ethernet",
        "6. Hạn chế và rủi ro còn lại",
        "7. Kế hoạch tuần tiếp theo",
        "Phụ lục: Tệp kết quả",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(entry)
        set_run_font(r, 12)
    doc.add_page_break()

    add_heading(doc, "1. Mục tiêu tuần", 1)
    add_paragraph(doc, "Mục tiêu của tuần là đưa mô hình QNN phát hiện cell và droplet từ mức kiểm thử trên ảnh/video sang một thiết kế có thể chạy thực tế trên FPGA Zybo Z7-10. Hệ thống sử dụng hai ROI liên tiếp: ROI 1 phát hiện ứng viên, ROI 2 xác nhận lại cùng đối tượng để mỗi cell hoặc droplet chỉ được đếm một lần khi đi qua vùng kiểm định.")
    add_paragraph(doc, "Các tiêu chí đánh giá gồm thông lượng tối thiểu 60 khung hình/giây cho hai ROI, kết quả QNN ổn định sau khi tăng tần số, timing đạt, tài nguyên không vượt khả năng XC7Z010 và video hiển thị đúng kết quả của từng frame.")

    add_heading(doc, "2. Công việc đã thực hiện", 1)
    add_paragraph(doc, "Mô hình QNN W4A6 đầu vào 96 x 96 được tích hợp vào PL của Zybo Z7-10. Thiết kế gồm giao tiếp DMA/UART sparse để kiểm tra phần cứng, hai lệnh suy luận ROI trên mỗi frame, bộ theo dõi ID và bộ ghép cặp sự kiện ROI 1 sang ROI 2.")
    add_paragraph(doc, "Tần số logic được tăng từ 100 MHz lên 104 MHz. Sau implementation, các tensor đầu ra của 16 giao dịch ROI trích từ video thực vẫn giống bit tuyệt đối so với baseline 100 MHz. Bản thử nghiệm 140 MHz trả tensor rỗng và checksum 00000000 nên đã bị loại khỏi kết quả.")
    doc.add_picture(str(architecture), width=Inches(6.6))
    add_caption(doc, "Hình 1. Luồng xử lý và cơ chế đếm xác nhận hai ROI.")

    add_heading(doc, "3. Kết quả triển khai QNN trên FPGA", 1)
    add_metric_table(doc, [
        ("Bo mạch và chip", "Digilent Zybo Z7-10, Xilinx XC7Z010-1CLG400C"),
        ("Mô hình", "QNN W4A6, ảnh xám ROI 96 x 96, hai ROI trên mỗi frame"),
        ("Tần số logic", "104 MHz"),
        ("Thông lượng đo trực tiếp trên FPGA", "121.23 ROI/s = 60.61 frame/s với hai ROI"),
        ("Baseline 100 MHz", "58.28 frame/s với hai ROI"),
        ("Độ trễ QNN trung bình", "8.586 ms / ROI ở 104 MHz"),
        ("Kiểm chứng đầu ra", "16 tensor video thực giống bit tuyệt đối với baseline; checksum CFC48C10"),
        ("Timing", "WNS +0.586 ns, WHS +0.027 ns; timing đạt"),
        ("Công suất ước tính", "1.872 W"),
    ])
    doc.add_picture(str(fps_chart), width=Inches(6.2))
    add_caption(doc, "Hình 2. Thông lượng hai ROI tăng từ 58.28 lên 60.61 FPS.")
    doc.add_picture(str(resource_chart), width=Inches(6.2))
    add_caption(doc, "Hình 3. Mức sử dụng tài nguyên sau implementation.")
    add_paragraph(doc, "Mức dùng LUT là 66.27%, FF 41.97%, BRAM 9.17% và DSP 13.75%. LUT là tài nguyên cần theo dõi sát nhất nếu mở rộng mô hình hoặc bổ sung khối xử lý video trong PL; BRAM và DSP hiện còn đủ khoảng trống cho buffer/tiền xử lý có chọn lọc.")

    add_heading(doc, "4. Kiểm chứng video và thuật toán đếm hai ROI", 1)
    add_paragraph(doc, "Video kiểm chứng sử dụng nguồn 3.4.mp4, xử lý 300 frame ở 30 FPS. Hai ROI có kích thước 120 x 120 pixel: ROI ứng viên [628, 410, 748, 530] và ROI xác nhận [748, 410, 868, 530]. Mỗi box và track được gắn với chính source frame đang hiển thị, vì vậy không dùng kết quả cũ để vẽ lên frame mới.")
    add_paragraph(doc, "Quy tắc đếm là: một đối tượng chỉ tăng bộ đếm khi đã được phát hiện ở ROI 1 và ghép cặp hợp lệ ở ROI 2. Hệ thống tạo 18 ứng viên cell và 14 ứng viên droplet; sau xác nhận thu được 16 sự kiện, gồm 8 cell và 8 droplet. Track đã xác nhận được giữ box dự đoán tối đa 10 frame mất quan sát để giảm bỏ sót khi đối tượng tạm chạm viền ROI.")
    doc.add_picture(str(preview), width=Inches(6.4))
    add_caption(doc, "Hình 4. Frame video đã đồng bộ kết quả QNN, hai ROI và tracker ID.")
    add_metric_table(doc, [
        ("Frame kiểm chứng", "300 frame từ video 3.4.mp4"),
        ("Kết quả đã xác nhận", "8 cell; 8 droplet; 16 sự kiện"),
        ("Cơ chế chống đếm đôi", "ID từ ROI 1 phải khớp ROI 2 trước khi tăng bộ đếm"),
        ("Đồng bộ hiển thị", "Box, track và frame_id thuộc cùng một source frame"),
        ("Tốc độ tạo video đồng bộ", "6.23 FPS do vòng UART đồng bộ từng frame; không phải tốc độ QNN trong PL"),
        ("Giới hạn UART cũ", "Khoảng 7.46 cập nhật hai ROI/s ở 2 Mbaud; không phù hợp để truyền video realtime"),
    ])

    add_heading(doc, "5. Tích hợp camera Phantom Ethernet", 1)
    add_paragraph(doc, "Bộ SDK camera đã được kiểm tra là Vision Research Phantom SDK 13.8.804.66. SDK cung cấp live image, discovery camera bằng broadcast/UDP và truyền ảnh qua Ethernet, nhưng chỉ có binary Windows x64. Không có thư viện Linux ARM hoặc mô tả giao thức mạng đủ để đưa SDK trực tiếp lên Cortex-A9 của Zybo.")
    add_paragraph(doc, "Kiến trúc triển khai khả thi trước mắt là PC dùng Phantom SDK chỉ để nhận frame thô từ camera; frame được chuyển bằng Gigabit Ethernet sang Zybo. Tại FPGA thực hiện crop ROI, lượng tử hóa, QNN, tracking và đếm; FPGA trả metadata gồm frame_id, box, ID, lớp và bộ đếm để PC vẽ live. Với Mono8 640 x 640 ở 60 FPS, payload khoảng 196.6 Mbit/s, nằm trong khả năng Gigabit Ethernet. UART chỉ giữ vai trò debug và nạp firmware.")

    add_heading(doc, "6. Hạn chế và rủi ro còn lại", 1)
    add_paragraph(doc, "Kết quả 60.61 FPS là benchmark của lõi QNN hai ROI trong FPGA, chưa phải tốc độ end-to-end từ camera. Để đạt 60 FPS thực tế, đường truyền frame phải chuyển từ UART sang Gigabit Ethernet hoặc AXI-Stream/VDMA; frame cần được gắn frame_id để tránh box trễ hoặc giật khi hiển thị.")
    add_paragraph(doc, "Chất lượng box và phân loại còn phụ thuộc trực tiếp vào nhãn dữ liệu. Các trường hợp cell/droplet sát viền, chồng lấp hoặc độ tương phản thấp cần được giữ trong tập validation/test độc lập. Không nên đánh giá mô hình chỉ bằng các frame đã dùng để chọn ngưỡng hoặc tinh chỉnh tracker.")

    add_heading(doc, "7. Kế hoạch tuần tiếp theo", 1)
    add_paragraph(doc, "Thứ nhất, kết nối camera Phantom, xác nhận model/IP và đo FPS live thực bằng SDK. Thứ hai, xây dựng cầu raw-frame Gigabit Ethernet giữa PC và Zybo với frame_id, CRC và bộ đệm vòng. Thứ ba, chuyển crop/resize/quantize hai ROI vào đường AXI/DMA để QNN nhận dữ liệu mà không cần gửi từng ROI qua UART.")
    add_paragraph(doc, "Song song, tiếp tục rà soát nhãn cell và droplet theo các trường hợp khó, huấn luyện lại mô hình trên tập dữ liệu đã duyệt và so sánh accuracy, recall ở biên ROI trước khi thay weight cho QNN. Cuối cùng, kiểm tra 60 FPS end-to-end trên camera thật và lưu log đếm dưới dạng CSV theo frame/timestamp.")

    add_heading(doc, "Phụ lục: Tệp kết quả", 1)
    add_paragraph(doc, "Bitstream 104 MHz: final_results/zybo_z7_10_qnn_60fps_104mhz/artifacts/zybo_qnn_wrapper.bit.")
    add_paragraph(doc, "Firmware sparse: final_results/zybo_z7_10_qnn_60fps_104mhz/artifacts/qnn_dma_uart_sparse.elf.")
    add_paragraph(doc, "Video kiểm chứng: final_results/zybo_z7_10_qnn_60fps_104mhz/video_3_4_exact_tracking/zybo_dual_roi_exact_arty_style.mp4.")
    add_paragraph(doc, "Sự kiện đã xác nhận: final_results/zybo_z7_10_qnn_60fps_104mhz/video_3_4_exact_tracking/confirmed_events.csv.")
    add_paragraph(doc, "Nghiên cứu kết nối camera: reports/phantom_camera_ethernet_sdk_research.md.")

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
